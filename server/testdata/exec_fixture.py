"""테스트용 임원회의 PPT 견본 — **실물에서 실측한 구조의 기록**.

실물 파일(USDMPD215 2026년_20260814_유니에버 수행전략회의_…)은 사내 자료라
저장소에 넣을 수 없다. 대신 그 파일에서 잰 값을 그대로 재현한 견본을 만든다.

    슬라이드 10.83 × 7.50 in
    로드맵 표 7행 × 18열
    병합  (0,0) 2×2 'Project' / (0,2) 1×12 '2026년' / (0,14~17) 2×1
          (2,0) 2×1 / (4,0) 3×1 / 데이터 행의 가로 병합(진행 구간)
    색    머리글 #FFFFCC, 단계 #FBE5D6 #DEEBF7 #E2F0D9 #FFFF00 #92D050
    진행 현황 / 향후 계획 = 3행 2열 표 **하나**
    이슈 리스트 = 2행 1열 표

**여기 적힌 숫자를 마음대로 바꾸지 말 것.** 이건 우리가 정한 양식이 아니라
임원들이 이미 쓰고 있는 양식이다. 바꾸면 변환기 테스트는 통과하고 실물은 깨진다.

바이너리를 저장소에 두지 않고 매번 만든다 — 견본이 낡아 실물과 어긋나는 일을 막는다.
"""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

SLIDE_W_IN = 10.83
SLIDE_H_IN = 7.50
ROWS, COLS = 7, 18
HEADER_BG = "FFFFCC"
STAGE = ["FBE5D6", "DEEBF7", "E2F0D9", "FFFF00", "92D050"]

# (행, 시작열, 칸수, 라벨, 색) — 진행 구간은 색칠이 아니라 **가로 병합 + 단계 이름**이다.
BARS = [
    (2, 2, 4, "요건정의", STAGE[0]),
    (2, 6, 5, "설계·구축", STAGE[1]),
    (3, 3, 3, "분석", STAGE[2]),
    (3, 6, 7, "개발", STAGE[1]),
    (4, 2, 7, "완전자율형에이전틱엔진, 설비이상탐지, 데이터", STAGE[3]),
    (5, 9, 5, "PoC", STAGE[4]),
    (6, 4, 6, "설계", STAGE[0]),
]
# 표의 진짜 크기는 열 너비·행 높이의 합이다(ext 가 아니라).
TRUE_TABLE_W_IN = 1.15 + 1.35 + 0.42 * 12 + 0.62 * 4     # 10.02in
STALE_EXT_IN = 2.5                                        # 파일에 적힌 낡은 값

EXPECTED_MERGES = [
    (0, 0, 2, 2), (0, 2, 1, 12), (0, 14, 2, 1), (0, 15, 2, 1),
    (0, 16, 2, 1), (0, 17, 2, 1), (2, 0, 2, 1), (4, 0, 3, 1),
]


def _fill(cell, hexstr: str) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(hexstr)


def build() -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    slide = prs.slides.add_slide(prs.slide_layouts[6])       # 빈 레이아웃

    tb = slide.shapes.add_textbox(Inches(0.35), Inches(0.2), Inches(7.0), Inches(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "2026년 수행전략회의 — 진행보고"
    run.font.size = Pt(20)
    run.font.bold = True

    gf = slide.shapes.add_table(ROWS, COLS, Inches(0.30), Inches(0.85),
                                Inches(10.2), Inches(3.1))
    tbl = gf.table
    tbl.columns[0].width = Emu(int(Inches(1.15)))
    tbl.columns[1].width = Emu(int(Inches(1.35)))
    for c in range(2, 14):
        tbl.columns[c].width = Emu(int(Inches(0.42)))
    for c in range(14, 18):
        tbl.columns[c].width = Emu(int(Inches(0.62)))

    tbl.cell(0, 0).merge(tbl.cell(1, 1))
    tbl.cell(0, 0).text = "Project"
    tbl.cell(0, 2).merge(tbl.cell(0, 13))
    tbl.cell(0, 2).text = "2026년"
    for c, label in ((14, "2027년"), (15, "계획 M/M"), (16, "투입 M/M"), (17, "비고(투입인원)")):
        tbl.cell(0, c).merge(tbl.cell(1, c))
        tbl.cell(0, c).text = label
    for i in range(12):
        tbl.cell(1, 2 + i).text = str(i + 1)
    for r in (0, 1):
        for c in range(COLS):
            _fill(tbl.cell(r, c), HEADER_BG)

    tbl.cell(2, 0).merge(tbl.cell(3, 0))
    tbl.cell(2, 0).text = "삼일"
    tbl.cell(4, 0).merge(tbl.cell(6, 0))
    tbl.cell(4, 0).text = "자세 솔루션 사업"
    for r, name in ((2, "회계 고도화"), (3, "세무 자동화"),
                    (4, "완전자율형 엔진"), (5, "설비 이상탐지"), (6, "데이터 플랫폼")):
        tbl.cell(r, 1).text = name

    for r, c0, span, label, color in BARS:
        if span > 1:
            tbl.cell(r, c0).merge(tbl.cell(r, c0 + span - 1))
        tbl.cell(r, c0).text = label
        _fill(tbl.cell(r, c0), color)

    for r in range(2, ROWS):
        tbl.cell(r, 15).text = "2.0"
        tbl.cell(r, 16).text = "1.5"
        tbl.cell(r, 17).text = "홍길동 외 2"

    # **일부러 낡은 크기를 심는다 — 반드시 표를 다 만든 뒤에.**
    # 열 너비를 지정하면 python-pptx 가 graphicFrame 의 ext 를 다시 계산해 덮어쓴다.
    # 그래서 먼저 심으면 흔적도 없이 사라진다(그렇게 해봤다가 테스트가 버그를 못 잡았다).
    #
    # 실물 임원회의 파일이 이 상태였다 — ext 가 갱신되지 않아 표 3개가 전부
    # 같은 작은 정사각형(315×315)으로 들어왔다. 위치는 맞았고 크기만 낡았다.
    gf.width = Emu(int(Inches(STALE_EXT_IN)))
    gf.height = Emu(int(Inches(STALE_EXT_IN)))

    t2 = slide.shapes.add_table(3, 2, Inches(0.30), Inches(4.15),
                                Inches(6.6), Inches(1.9)).table
    t2.cell(0, 0).text = "진행 현황"
    t2.cell(0, 1).text = "향후 계획"
    t2.cell(1, 0).text = "· 회계 고도화 요건정의 완료"
    t2.cell(1, 1).text = "· 설계 착수"
    t2.cell(2, 0).text = "· PoC 환경 구성"
    t2.cell(2, 1).text = "· 사용자 검증"

    t3 = slide.shapes.add_table(2, 1, Inches(7.10), Inches(4.15),
                                Inches(3.4), Inches(1.9)).table
    t3.cell(0, 0).text = "이슈 리스트"
    t3.cell(1, 0).text = "· 인력 충원 지연"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_many(n: int) -> bytes:
    """슬라이드 n장짜리 — 상한 검사용."""
    prs = Presentation()
    for _ in range(n):
        prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    import pathlib
    out = pathlib.Path(__file__).with_name("exec_roadmap.pptx")
    out.write_bytes(build())
    print("wrote", out)
