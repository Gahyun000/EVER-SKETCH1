"""실물 PPT 업로드 → 슬라이드별 배부.

지키려는 것
  1. 관리자만 올리고 배부한다 (작성자·열람자·비로그인은 전부 막힌다)
  2. 슬라이드가 담당자별로 나뉘고, 남의 슬라이드는 안 간다
  3. 두 번 눌러도 두 장이 생기지 않는다
  4. 이미 배부한 회차의 원본은 바뀌지 않는다 (바뀌면 나간 자료와 어긋난다)
  5. 한 사람의 편집이 다른 사람 문서로 번지지 않는다 (얕은 복사 사고)
"""
import os
import pathlib
import tempfile

_tmp = tempfile.mkdtemp()
os.environ["EVER_SKETCH_DB"] = str(pathlib.Path(_tmp) / "decks.db")

import pytest  # noqa: E402
from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from server import auth as auth_store  # noqa: E402
from server import cycle_decks as decks_store  # noqa: E402
from server import cycles as cycles_store  # noqa: E402
from server import projects as projects_store  # noqa: E402
from server.routes_auth import router as auth_router  # noqa: E402
from server.routes_cycles import router as cycles_router  # noqa: E402
from server.routes_projects import router as proj_router  # noqa: E402
from server.testdata import exec_fixture as fx  # noqa: E402

_DB = str(pathlib.Path(_tmp) / "decks.db")


def make_app() -> FastAPI:
    app = FastAPI()
    app.include_router(auth_router)
    app.include_router(proj_router)
    app.include_router(cycles_router)
    return app


def _pptx_bytes(slides: int = 3) -> bytes:
    """표지 + 임원별 로드맵 2장짜리 견본. (실측 구조는 exec_fixture 참조)"""
    import io as _io

    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation(_io.BytesIO(fx.build()))
    while len(prs.slides) < slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = "슬라이드 %d" % len(prs.slides)
        run.font.size = Pt(24)
    buf = _io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def deck_bytes() -> bytes:
    return _pptx_bytes(3)


@pytest.fixture()
def ctx():
    os.environ["EVER_SKETCH_DB"] = _DB
    for suffix in ("", "-wal", "-shm"):
        f = pathlib.Path(_DB + suffix)
        if f.exists():
            f.unlink()

    auth_store.ensure_seed_admin("adminpw12345")
    admin = [u for u in auth_store.list_users() if u["login_id"] == "admin"][0]

    def mk(login, role, name, dept="사업본부"):
        u = auth_store.signup(login, "password123", name, dept, "writer")
        if role:
            auth_store.approve(admin["id"], u["id"], role)
        return auth_store.get_user(u["id"])

    execs = [mk("exec%d" % i, "writer", "임원%d" % i, "제%d본부" % i) for i in (1, 2)]
    viewer = mk("viewer", "viewer", "열람자")

    def as_user(login, pw="password123"):
        c = TestClient(make_app())
        r = c.post("/api/auth/login", json={"login_id": login, "password": pw})
        assert r.status_code == 200, r.text
        return c

    admin_client = as_user("admin", "adminpw12345")
    cid = admin_client.post("/api/cycles", json={"period_ym": "2026-10"}).json()["cycle"]["id"]

    return {"admin": admin, "execs": execs, "viewer": viewer, "cid": cid,
            "as_user": as_user, "as_admin": lambda: as_user("admin", "adminpw12345"),
            "anon": TestClient(make_app())}


def _upload(client, cid, data, name="exec.pptx"):
    return client.post("/api/cycles/%s/deck" % cid,
                       files={"file": (name, data,
                                       "application/vnd.openxmlformats-officedocument."
                                       "presentationml.presentation")})


# ══════════ 권한 ══════════
def test_작성자는_올릴_수_없다(ctx, deck_bytes):
    r = _upload(ctx["as_user"]("exec1"), ctx["cid"], deck_bytes)
    assert r.status_code == 403


def test_열람자는_올릴_수_없다(ctx, deck_bytes):
    assert _upload(ctx["as_user"]("viewer"), ctx["cid"], deck_bytes).status_code == 403


def test_비로그인은_올릴_수_없다(ctx, deck_bytes):
    assert _upload(ctx["anon"], ctx["cid"], deck_bytes).status_code == 401


def test_작성자는_슬라이드_목록을_볼_수_없다(ctx, deck_bytes):
    _upload(ctx["as_admin"](), ctx["cid"], deck_bytes)
    r = ctx["as_user"]("exec1").get("/api/cycles/%s/deck" % ctx["cid"])
    assert r.status_code == 403, "작성자에게 전체 슬라이드가 보이면 남의 자료가 새어 나간다."


def test_작성자는_배부할_수_없다(ctx, deck_bytes):
    _upload(ctx["as_admin"](), ctx["cid"], deck_bytes)
    r = ctx["as_user"]("exec1").post(
        "/api/cycles/%s/deck/distribute" % ctx["cid"],
        json={"assignments": [{"slide": 0, "user_id": ctx["execs"][0]["id"]}]})
    assert r.status_code == 403


# ══════════ 업로드 ══════════
def test_업로드하면_슬라이드_목록이_돌아온다(ctx, deck_bytes):
    r = _upload(ctx["as_admin"](), ctx["cid"], deck_bytes)
    assert r.status_code == 200, r.text
    deck = r.json()["deck"]
    assert deck["slide_count"] == 3
    assert len(deck["slides"]) == 3
    assert deck["slides"][0]["tables"] == 3          # 실물 구조가 살아 있다
    assert "pages" not in deck, "목록 응답에 변환 결과를 실으면 화면이 뜨지 않는다."


def test_pptx가_아니면_거부한다(ctx):
    r = _upload(ctx["as_admin"](), ctx["cid"], b"not a pptx at all", "가짜.pptx")
    assert r.status_code == 400
    assert "PowerPoint" in r.json()["detail"]


def test_다시_올리면_갈아탄다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes, "첫번째.pptx")
    _upload(c, ctx["cid"], deck_bytes, "두번째.pptx")
    assert decks_store.get_deck(ctx["cid"])["filename"] == "두번째.pptx"
    # 두 벌이 남으면 어느 쪽을 배부했는지 알 수 없다
    assert len(decks_store.get_deck_pages(ctx["cid"])) == 3


def test_배부한_뒤에는_원본을_못_바꾼다(ctx, deck_bytes):
    """이미 나간 자료와 원본이 어긋나면, 취합 때 무엇이 정본인지 아무도 모른다."""
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    c.post("/api/cycles/%s/deck/distribute" % ctx["cid"],
           json={"assignments": [{"slide": 1, "user_id": ctx["execs"][0]["id"]}]})
    r = _upload(c, ctx["cid"], deck_bytes, "새파일.pptx")
    assert r.status_code == 400
    assert "새 회차" in r.json()["detail"]


def test_삭제하면_사라진다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    assert c.delete("/api/cycles/%s/deck" % ctx["cid"]).status_code == 200
    assert c.get("/api/cycles/%s/deck" % ctx["cid"]).status_code == 404


# ══════════ 배부 ══════════
def test_슬라이드가_담당자별로_나뉜다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    a, b = ctx["execs"]
    r = c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json={
        "assignments": [{"slide": 1, "user_id": a["id"]},
                        {"slide": 2, "user_id": b["id"]}],
    })
    assert r.status_code == 200, r.text
    assert r.json()["created_count"] == 2

    rows = cycles_store.list_cycle_projects(ctx["cid"])
    assert {p["owner_id"] for p in rows} == {a["id"], b["id"]}
    for p in rows:
        assert p["page_count"] == 1, "담당분만 가야 한다 — 3장이면 남의 슬라이드까지 갔다."


def test_공통_슬라이드는_모두에게_앞에_붙는다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    a, b = ctx["execs"]
    c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json={
        "assignments": [{"slide": 1, "user_id": a["id"]},
                        {"slide": 2, "user_id": b["id"]}],
        "common": [0],
    })
    for p in cycles_store.list_cycle_projects(ctx["cid"]):
        assert p["page_count"] == 2
        state = projects_store.get_project(p["id"])["state"]
        assert [pg["id"] for pg in state["pages"]] == [1, 2], "페이지 번호를 1부터 다시 매겨야 한다."
        # 표지(0번)에는 표가 3개 있다 — 맨 앞에 왔는지 내용으로 확인
        assert len([e for e in state["pages"][0]["els"] if e["type"] == "table"]) == 3


def test_공통_슬라이드는_중복되지_않는다(ctx, deck_bytes):
    """담당 슬라이드가 공통에도 들어 있으면 같은 장이 두 번 들어간다."""
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    a = ctx["execs"][0]
    c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json={
        "assignments": [{"slide": 0, "user_id": a["id"]}], "common": [0]})
    p = cycles_store.list_cycle_projects(ctx["cid"])[0]
    assert p["page_count"] == 1


def test_한_사람의_편집이_다른_사람에게_번지지_않는다(ctx, deck_bytes):
    """같은 슬라이드를 둘에게 주면 각자 자기 사본을 가져야 한다."""
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    a, b = ctx["execs"]
    c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json={
        "assignments": [{"slide": 0, "user_id": a["id"]},
                        {"slide": 0, "user_id": b["id"]}]})
    rows = cycles_store.list_cycle_projects(ctx["cid"])
    sa, sb = [projects_store.get_project(p["id"])["state"] for p in rows]
    ea = sa["pages"][0]["els"][0]
    eb = sb["pages"][0]["els"][0]
    assert ea is not eb
    ea["text"] = "내가 고친 글"
    assert eb["text"] != "내가 고친 글"


def test_배부는_멱등하다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    a = ctx["execs"][0]
    body = {"assignments": [{"slide": 1, "user_id": a["id"]}]}
    first = c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json=body).json()
    second = c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json=body).json()
    assert first["created_count"] == 1 and second["created_count"] == 0
    assert second["skipped_count"] == 1
    assert len(cycles_store.list_cycle_projects(ctx["cid"])) == 1


def test_배부하면_회차가_작성중이_된다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    assert cycles_store.get_cycle(ctx["cid"])["status"] == "draft"
    c.post("/api/cycles/%s/deck/distribute" % ctx["cid"],
           json={"assignments": [{"slide": 0, "user_id": ctx["execs"][0]["id"]}]})
    assert cycles_store.get_cycle(ctx["cid"])["status"] == "writing"


def test_배부본은_같은_회차_동료까지_열린다(ctx, deck_bytes):
    """'본인만' 이던 규칙을 바꿨다 — 같은 회의를 준비하는 사람끼리는 서로 본다.
    고치는 것은 여전히 본인 것만이다."""
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    a, b = ctx["execs"]
    c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json={
        "assignments": [{"slide": 1, "user_id": a["id"]},
                        {"slide": 2, "user_id": b["id"]}]})
    rows = cycles_store.list_cycle_projects(ctx["cid"])
    mine = [p for p in rows if p["owner_id"] == a["id"]][0]
    other = [p for p in rows if p["owner_id"] == b["id"]][0]
    c1 = ctx["as_user"]("exec1")
    assert c1.get("/api/projects/%s" % mine["id"]).status_code == 200
    assert c1.get("/api/projects/%s" % other["id"]).status_code == 200
    assert c1.put("/api/projects/%s" % other["id"],
                  json={"state": {"pages": []}, "name": "남의 장"}).status_code == 403


# ══════════ 잘못된 입력 ══════════
def test_없는_슬라이드_번호는_거부한다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    r = c.post("/api/cycles/%s/deck/distribute" % ctx["cid"],
               json={"assignments": [{"slide": 99, "user_id": ctx["execs"][0]["id"]}]})
    assert r.status_code == 400


def test_승인되지_않은_계정에는_배부하지_않는다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    r = c.post("/api/cycles/%s/deck/distribute" % ctx["cid"],
               json={"assignments": [{"slide": 0, "user_id": "없는사람"}]})
    assert r.status_code == 400


def test_올린_PPT가_없으면_배부할_수_없다(ctx):
    c = ctx["as_admin"]()
    r = c.post("/api/cycles/%s/deck/distribute" % ctx["cid"],
               json={"assignments": [{"slide": 0, "user_id": ctx["execs"][0]["id"]}]})
    assert r.status_code == 400
    assert "PPT" in r.json()["detail"]


def test_담당자를_아무도_지정하지_않으면_거부한다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    r = c.post("/api/cycles/%s/deck/distribute" % ctx["cid"], json={"assignments": []})
    assert r.status_code == 400


# ══════════ 배부 전 미리보기 ══════════
def test_미리보기는_표준_양식_한_장을_돌려준다(ctx):
    """배부 전에 눈으로 확인하게 한다 — 좌표계가 어긋나 빈 종이가 나간 적이 있다."""
    r = ctx["as_admin"]().get("/api/cycles/%s/preview" % ctx["cid"])
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["mode"] == "template"
    assert body["page"]["cardKey"] == "slide"
    tables = [e for e in body["page"]["els"] if e["type"] == "table"]
    assert len(tables) == 3


def test_미리보기가_올린_PPT의_그_장을_돌려준다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    r = c.get("/api/cycles/%s/preview?slide=0" % ctx["cid"])
    assert r.status_code == 200, r.text
    assert r.json()["mode"] == "deck"
    assert len([e for e in r.json()["page"]["els"] if e["type"] == "table"]) == 3


def test_없는_슬라이드_미리보기는_거부한다(ctx, deck_bytes):
    c = ctx["as_admin"]()
    _upload(c, ctx["cid"], deck_bytes)
    assert c.get("/api/cycles/%s/preview?slide=99" % ctx["cid"]).status_code == 400


def test_PPT를_안_올렸는데_슬라이드_미리보기를_요청하면_404(ctx):
    assert ctx["as_admin"]().get(
        "/api/cycles/%s/preview?slide=0" % ctx["cid"]).status_code == 404


def test_작성자는_미리보기를_볼_수_없다(ctx):
    """전 슬라이드가 담길 수 있는 창구다 — 남의 자료가 새면 안 된다."""
    assert ctx["as_user"]("exec1").get(
        "/api/cycles/%s/preview" % ctx["cid"]).status_code == 403


# ══════════ 서버에 모듈이 없을 때 ══════════
def test_모듈이_없으면_503과_고치는_법을_돌려준다(ctx, deck_bytes, monkeypatch):
    """운영 중에 실제로 난 사고를 화면 문구까지 못 박는다.

    python-pptx 가 venv 에 없어서 업로드가 죽었는데, 화면에는 'HTTP 500' 만 떴다.
    사용자는 자기 파일이 잘못된 줄 알고 몇 번이나 다시 저장해 올렸다.
    500(원인 불명)도 400(당신 파일 잘못)도 아닌 503(서버가 덜 갖춰짐)이어야 하고,
    무엇을 설치하면 되는지까지 문구에 있어야 한다.
    """
    import builtins

    real_import = builtins.__import__

    def no_pptx(name, *a, **kw):
        if name == "pptx" or name.startswith("pptx."):
            raise ImportError("No module named 'pptx'")
        return real_import(name, *a, **kw)

    monkeypatch.setattr(builtins, "__import__", no_pptx)
    r = _upload(ctx["as_admin"](), ctx["cid"], deck_bytes)
    assert r.status_code == 503, r.status_code
    detail = r.json()["detail"]
    assert "python-pptx" in detail
    assert "pip install" in detail


# ══════════ 변환기 판 · 낡은 자료 표시 ══════════
def test_올린_자료에_변환기_판이_남는다(ctx, deck_bytes):
    from server import pptx_import

    _upload(ctx["as_admin"](), ctx["cid"], deck_bytes)
    deck = decks_store.get_deck(ctx["cid"])
    assert deck["converter"] == pptx_import.CONVERTER_VERSION
    assert deck["stale"] is False


def test_예전_판으로_읽은_자료는_낡음으로_표시된다(ctx, deck_bytes):
    """이게 없어서 같은 확인을 세 번 반복했다.

    표 크기·표지 배경·제목 뽑기를 고쳤는데 화면에는 예전에 변환된 자료가 그대로
    떠 있었다. 고쳤다고 말했는데 사용자 눈에는 똑같이 보였고, 무엇이 잘못됐는지
    사용자도 나도 알 수 없었다. 판을 남겨 두면 화면이 그 사실을 말해 준다.
    """
    import sqlite3

    _upload(ctx["as_admin"](), ctx["cid"], deck_bytes)
    conn = sqlite3.connect(decks_store.cycles_store.projects_store._db_path())
    try:
        conn.execute("UPDATE CycleDecks SET converter='v0.9' WHERE cycle_id=?", (ctx["cid"],))
        conn.commit()
    finally:
        conn.close()

    deck = decks_store.get_deck(ctx["cid"])
    assert deck["stale"] is True
    assert deck["converter"] == "v0.9"


def test_판이_비어_있는_옛_자료도_낡음이다(ctx, deck_bytes):
    """판을 기록하기 전에 올린 자료 — 화면에서 다시 올리라고 알려야 한다."""
    import sqlite3

    _upload(ctx["as_admin"](), ctx["cid"], deck_bytes)
    conn = sqlite3.connect(decks_store.cycles_store.projects_store._db_path())
    try:
        conn.execute("UPDATE CycleDecks SET converter='' WHERE cycle_id=?", (ctx["cid"],))
        conn.commit()
    finally:
        conn.close()
    assert decks_store.get_deck(ctx["cid"])["stale"] is True


def test_변환_결과가_바뀌면_판을_올려야_한다():
    """판을 안 올리고 변환기를 고치면, 낡은 자료가 최신인 척한다.

    사람이 기억할 일이 아니라 목록이 기억할 일이다 — 판마다 무엇이 바뀌었는지
    pptx_import.py 주석에 적혀 있어야 한다.
    """
    import re
    from pathlib import Path

    src = Path(decks_store.__file__).with_name("pptx_import.py").read_text(encoding="utf-8")
    from server import pptx_import

    ver = pptx_import.CONVERTER_VERSION
    assert re.search(r"^#\s+%s\s+\S" % re.escape(ver), src, re.M), \
        "CONVERTER_VERSION=%s 의 변경 내용이 주석 목록에 없습니다." % ver
